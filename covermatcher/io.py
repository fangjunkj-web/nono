from pathlib import Path
import csv
from openpyxl import load_workbook

def load_topics(path):
    p=Path(path)
    rows=[]
    if p.suffix.lower()=='.csv':
        with p.open(encoding='utf-8-sig', newline='') as f:
            data=list(csv.reader(f))
    else:
        ws=load_workbook(p, read_only=True, data_only=True).active
        data=[[c for c in r] for r in ws.iter_rows(values_only=True)]
    if not data: return []
    header=[str(x or '').strip().lower() for x in data[0]]
    idx=next((i for i,x in enumerate(header) if x in {'title','topic','article title','文章标题','标题'}),0)
    for r in data[1:]:
        if idx < len(r) and r[idx]: rows.append(str(r[idx]).strip())
    return rows

def list_images(folder):
    exts={'.jpg','.jpeg','.png','.webp','.bmp','.tif','.tiff'}
    return [p for p in Path(folder).rglob('*') if p.is_file() and p.suffix.lower() in exts]
